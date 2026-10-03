"""Build a journal-specific Sexually Transmitted Infections submission package."""

from __future__ import annotations

import csv
import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import build_manuscript as bm  # noqa: E402


STI_ROOT = ROOT / "journal_submissions" / "STI"
SOURCE = STI_ROOT / "working" / "sti_manuscript_source.md"
FILES = STI_ROOT / "files"
ARCHIVE = STI_ROOT / "STI_Submission_Package.zip"
MAIN_MD = FILES / "01_Main_Manuscript" / "STI_Main_Manuscript.md"
MAIN_DOCX = FILES / "01_Main_Manuscript" / "STI_Main_Manuscript.docx"


TABLES = {
    "TABLE_1": bm.TABLES["TABLE_1"],
    "TABLE_2": {
        **bm.TABLES["TABLE_4"],
        "caption": "Table 2. Common-city mapped population and organisational change across economic periods",
    },
}

FIGURES = {
    "FIGURE_1": bm.FIGURES["FIGURE_1"],
    "FIGURE_2": {
        **bm.FIGURES["FIGURE_3"],
        "caption": "Figure 2. Common-city mapped population change during economic deterioration and recovery. The combined population expanded rapidly before 2011 and then grew more slowly, approaching a plateau through 2016-17.",
    },
}


def word_count(text: str) -> int:
    text = re.sub(r"(?m)^#+\s+.*$", " ", text)
    text = re.sub(r"\[\[(?:TABLE|FIGURE)_\d+\]\]", " ", text)
    text = re.sub(r"\[@[^\]]+\]", " ", text)
    return len(re.findall(r"\b[\w'-]+\b", text))


def counts(source: str, order: list[str]) -> dict[str, int]:
    abstract = source.split("## Abstract", 1)[1].split("## Key messages", 1)[0]
    main = source.split("## Introduction", 1)[1].split("## Declarations", 1)[0]
    return {
        "abstract_words": word_count(abstract),
        "main_text_words": word_count(main),
        "references": len(order),
        "tables": len(TABLES),
        "figures": len(FIGURES),
    }


def configure_plain(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(5)


def add_centered_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    bm.add_inline_runs(p, bm.TITLE, size=16)
    for run in p.runs:
        run.bold = True


def create_title_page(path: Path, metrics: dict[str, int]) -> None:
    doc = Document()
    configure_plain(doc)
    add_centered_title(doc)
    bm.add_author_block(doc, font_size=10.5, include_ssrn=True, wrap_corresponding=True)
    items = [
        "Article type: Original research",
        f"Main-text word count: {metrics['main_text_words']}",
        f"Abstract word count: {metrics['abstract_words']}",
        f"Display items: {metrics['tables']} tables and {metrics['figures']} figures (4 total)",
        f"References: {metrics['references']}",
        "Keywords: female sex workers; HIV surveillance; economic conditions; Pakistan; population size estimation",
        "Funding: No specific grant was received.",
        "Competing interests: The authors declare no competing interests.",
    ]
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        bm.add_inline_runs(p, item, size=10.5)
    doc.core_properties.title = f"Title page - {bm.TITLE}"
    doc.core_properties.author = "Qazi Raza; Amna Khan; Maham Qazi"
    doc.save(path)


def create_cover_letter(path: Path) -> None:
    doc = Document()
    configure_plain(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    bm.add_inline_runs(p, "3 October 2026", size=11)
    for line in ("Editor", "Sexually Transmitted Infections", "BMJ Publishing Group"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        bm.add_inline_runs(p, line, size=11)
    doc.add_paragraph()
    p = doc.add_paragraph()
    bm.add_inline_runs(p, "Dear Editor,", size=11)
    paragraphs = [
        f"Please consider our Original Research manuscript, \"{bm.TITLE},\" for publication in Sexually Transmitted Infections.",
        "The study links Pakistan's complete published female sex work mapping series to contemporaneous district and national economic indicators. Its central finding is that economic deterioration coincided with rapid growth and reorganisation of the mapped population, whereas improved conditions coincided with marked growth deceleration toward a plateau. A cross-sectional density measure alone did not capture this relationship. This distinction is directly relevant to HIV surveillance, prevention planning, and outreach to populations moving into less visible or intermediary-linked settings.",
        "The paper fits the journal's epidemiological, sociological, and HIV remit and extends Pakistani surveillance studies previously published in STI. The submitted version complies with the Original Research limits: no more than 3,000 main-text words, a structured abstract under 300 words, four combined tables and figures, 30 references, and the required key-messages box. A completed STROBE checklist and detailed supplementary appendix accompany the manuscript.",
        "This manuscript is original, has not been published, and is not under consideration elsewhere. All authors made the stated contributions, approved the final manuscript, and approve submission to STI. The authors declare no competing interests and received no specific funding. The study used only public aggregate reports and de-identified secondary data; ethical approval and individual informed consent were not required. Derived data, code, and validation materials are available at https://github.com/QaziRaza/FSW.",
        "Thank you for considering this work.",
    ]
    for text in paragraphs:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        bm.add_inline_runs(p, text, size=11)
    for line in (
        "Sincerely,",
        "Qazi Raza, MPH",
        "Health Services Academy, NIH Complex, Park Road, Chak Shahzad, Islamabad, 44000, Pakistan",
        "Email: qaziraza.rq@gmail.com",
        "ORCID: https://orcid.org/0000-0003-4303-3390",
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        bm.add_inline_runs(p, line, size=11)
    doc.core_properties.title = "Cover letter to Sexually Transmitted Infections"
    doc.core_properties.author = "Qazi Raza"
    doc.save(path)


def write_table_csv(path: Path, spec: dict) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(spec["headers"])
        writer.writerows(spec["rows"])


def write_readme(metrics: dict[str, int]) -> None:
    text = f"""# Sexually Transmitted Infections submission package

Article type: Original research

Study: {bm.TITLE}

## Verified journal limits

- Main text: {metrics['main_text_words']} words (maximum 3,000)
- Structured abstract: {metrics['abstract_words']} words (maximum 300)
- Tables and figures: {metrics['tables']} tables plus {metrics['figures']} figures = 4 total (maximum 4)
- References: {metrics['references']} (maximum 30)
- Key messages: present under all three required headings
- Reporting checklist: completed STROBE cross-sectional checklist included
- Data availability statement: present, with public repository URL

## Files

- `01_Main_Manuscript/` - line-numbered STI-formatted manuscript and inspectable Markdown
- `02_Title_Page/` - separate title page with author details and verified counts
- `03_Figures/` - the two main figures as PNG and 600 dpi TIFF files
- `04_Tables/` - editable CSV versions of the two main tables
- `05_Supplementary_Material/` - detailed analytical appendix and reproducibility archive
- `06_Reporting_Checklist/` - completed STROBE checklist
- `07_Cover_Letter/` - journal-specific cover letter

The comprehensive master manuscript remains outside this folder. This directory contains the distinct STI submission version.
"""
    (FILES / "README.md").write_text(text, encoding="utf-8")


def make_archive() -> None:
    with zipfile.ZipFile(ARCHIVE, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for path in sorted(p for p in FILES.rglob("*") if p.is_file()):
            archive.write(path, path.relative_to(FILES))


def build() -> None:
    if FILES.exists():
        resolved = FILES.resolve()
        if resolved.parent != STI_ROOT.resolve() or resolved.name != "files":
            raise RuntimeError(f"Refusing to reset unexpected path: {resolved}")
        shutil.rmtree(resolved)
    for relative in (
        "01_Main_Manuscript",
        "02_Title_Page",
        "03_Figures",
        "04_Tables",
        "05_Supplementary_Material",
        "06_Reporting_Checklist",
        "07_Cover_Letter",
    ):
        (FILES / relative).mkdir(parents=True, exist_ok=True)

    source = SOURCE.read_text(encoding="utf-8")
    resolved, order = bm.resolve_citations(source)
    metrics = counts(source, order)
    if metrics["main_text_words"] > 3000:
        raise AssertionError(f"STI main text exceeds 3,000 words: {metrics['main_text_words']}")
    if metrics["abstract_words"] > 300:
        raise AssertionError(f"STI abstract exceeds 300 words: {metrics['abstract_words']}")
    if metrics["references"] > 30:
        raise AssertionError(f"STI reference count exceeds 30: {metrics['references']}")
    if metrics["tables"] + metrics["figures"] > 4:
        raise AssertionError("STI display-item count exceeds four")

    bm.SOURCE = SOURCE
    bm.MARKDOWN_OUT = MAIN_MD
    bm.DOCX_OUT = MAIN_DOCX
    bm.TABLES = TABLES
    bm.FIGURES = FIGURES
    bm.produce_markdown(resolved, order)
    bm.build_docx(resolved, order)

    create_title_page(FILES / "02_Title_Page" / "STI_Title_Page.docx", metrics)
    create_cover_letter(FILES / "07_Cover_Letter" / "STI_Cover_Letter.docx")

    figures = FILES / "03_Figures"
    source_figures = ROOT / "submission_packages" / "files" / "04_Figures"
    for source_name, target_name in (
        ("Figure_1_Recent_economic_deterioration_and_mapped_FSW_density.png", "STI_Figure_1_Deterioration_and_FSW_density.png"),
        ("Figure_1_Recent_economic_deterioration_and_mapped_FSW_density_600dpi.tiff", "STI_Figure_1_Deterioration_and_FSW_density_600dpi.tiff"),
        ("Figure_3_Mapped_population_change_during_deterioration_and_recovery.png", "STI_Figure_2_Deterioration_and_recovery.png"),
        ("Figure_3_Mapped_population_change_during_deterioration_and_recovery_600dpi.tiff", "STI_Figure_2_Deterioration_and_recovery_600dpi.tiff"),
    ):
        shutil.copy2(source_figures / source_name, figures / target_name)

    write_table_csv(FILES / "04_Tables" / "STI_Table_1_Surveillance_rounds.csv", TABLES["TABLE_1"])
    write_table_csv(FILES / "04_Tables" / "STI_Table_2_Common_city_change.csv", TABLES["TABLE_2"])

    shutil.copy2(
        ROOT / "submission_packages" / "files" / "05_Supplementary_Material" / "Paper1_Supplementary_Appendix.docx",
        FILES / "05_Supplementary_Material" / "STI_Supplementary_Appendix.docx",
    )
    shutil.copy2(
        ROOT / "submission_packages" / "Paper1_Reproducibility_Data_and_Code.zip",
        FILES / "05_Supplementary_Material" / "STI_Reproducibility_Data_and_Code.zip",
    )
    shutil.copy2(
        ROOT / "submission_packages" / "files" / "06_Reporting_Checklist" / "Paper1_STROBE_Cross_Sectional_Checklist.docx",
        FILES / "06_Reporting_Checklist" / "STI_STROBE_Cross_Sectional_Checklist.docx",
    )
    (FILES / "01_Main_Manuscript" / "STI_Vancouver_References.txt").write_text(
        "\n".join(f"{i}. {bm.REFERENCES[key]}" for i, key in enumerate(order, start=1)) + "\n",
        encoding="utf-8",
    )
    write_readme(metrics)
    make_archive()
    print(metrics)
    print(f"Built {FILES}")
    print(f"Built {ARCHIVE}")


if __name__ == "__main__":
    build()
