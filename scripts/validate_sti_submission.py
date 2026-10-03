"""Validate the journal-specific STI submission package."""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_sti_submission import word_count  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
STI = ROOT / "journal_submissions" / "STI"
FILES = STI / "files"
TITLE = "Economic conditions and the changing scale and organisation of female sex work in Pakistan: a multi-period ecological study"
checks: list[tuple[str, bool]] = []


def check(name: str, condition: bool) -> None:
    checks.append((name, bool(condition)))


main_md = (FILES / "01_Main_Manuscript" / "STI_Main_Manuscript.md").read_text(encoding="utf-8")
source = (STI / "working" / "sti_manuscript_source.md").read_text(encoding="utf-8")
main_doc = Document(FILES / "01_Main_Manuscript" / "STI_Main_Manuscript.docx")
main_doc_text = "\n".join(p.text for p in main_doc.paragraphs)
abstract = source.split("## Abstract", 1)[1].split("## Key messages", 1)[0]
main_text = source.split("## Introduction", 1)[1].split("## Declarations", 1)[0]
reference_text = main_md.split("## References", 1)[1]
reference_numbers = re.findall(r"(?m)^(\d+)\. ", reference_text)

check("approved title present", TITLE in main_doc_text)
check("main text no more than 3000 words", word_count(main_text) <= 3000)
check("structured abstract no more than 300 words", word_count(abstract) <= 300)
for heading in ("Objectives", "Methods", "Results", "Conclusions"):
    check(f"abstract heading: {heading}", f"### {heading}" in abstract)
for heading in (
    "What is already known on this topic",
    "What this study adds",
    "How this study might affect research, practice or policy",
):
    check(f"key-message heading: {heading}", f"### {heading}" in main_md)
check("no more than 30 references", len(reference_numbers) <= 30)
check("four display items", len(main_doc.tables) == 2 and len(main_doc.inline_shapes) == 2)
check("data availability statement present", "https://github.com/QaziRaza/FSW" in main_doc_text)
check("final ethics statement present", "Ethical approval and individual informed consent were not required" in main_doc_text)
check("final author approval present", "All authors read and approved the final manuscript" in main_doc_text)
check("no unresolved citation keys", "[@" not in main_md)
check("no unresolved build markers", "[[" not in main_md and "]]" not in main_md)

required = [
    "01_Main_Manuscript/STI_Main_Manuscript.docx",
    "02_Title_Page/STI_Title_Page.docx",
    "03_Figures/STI_Figure_1_Deterioration_and_FSW_density_600dpi.tiff",
    "03_Figures/STI_Figure_2_Deterioration_and_recovery_600dpi.tiff",
    "04_Tables/STI_Table_1_Surveillance_rounds.csv",
    "04_Tables/STI_Table_2_Common_city_change.csv",
    "05_Supplementary_Material/STI_Supplementary_Appendix.docx",
    "05_Supplementary_Material/STI_Reproducibility_Data_and_Code.zip",
    "06_Reporting_Checklist/STI_STROBE_Cross_Sectional_Checklist.docx",
    "07_Cover_Letter/STI_Cover_Letter.docx",
    "README.md",
]
for relative in required:
    path = FILES / relative
    check(f"required file: {relative}", path.is_file() and path.stat().st_size > 0)

archive = STI / "STI_Submission_Package.zip"
check("submission archive exists", archive.is_file() and archive.stat().st_size > 0)
if archive.is_file():
    with zipfile.ZipFile(archive) as handle:
        check("submission archive integrity", handle.testzip() is None)

failures = [name for name, passed in checks if not passed]
for name, passed in checks:
    print(f"{'PASS' if passed else 'FAIL'} | {name}")
print(f"SUMMARY | {len(checks) - len(failures)}/{len(checks)} checks passed")
print(f"MAIN TEXT WORDS | {word_count(main_text)}")
print(f"ABSTRACT WORDS | {word_count(abstract)}")
print(f"REFERENCES | {len(reference_numbers)}")
print(f"TABLES | {len(main_doc.tables)}")
print(f"FIGURES | {len(main_doc.inline_shapes)}")
sys.exit(1 if failures else 0)
