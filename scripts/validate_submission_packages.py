"""Validate the Paper 1 journal submission archives."""

from pathlib import Path
import csv
import re
import sys
import zipfile

from docx import Document
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "submission_packages"
FILES = OUT / "files"
TITLE = "Economic conditions and the changing scale and organisation of female sex work in Pakistan: a multi-period ecological study"
checks = []


def check(label, condition):
    checks.append((label, bool(condition)))


expected_zips = [
    "Paper1_Tables_and_Figures.zip",
    "Paper1_Supplementary_Materials.zip",
    "Paper1_Reproducibility_Data_and_Code.zip",
    "Paper1_References.zip",
    "Paper1_Complete_Submission_Materials.zip",
]
for name in expected_zips:
    path = OUT / name
    check(f"archive exists: {name}", path.is_file() and path.stat().st_size > 0)
    if path.is_file():
        with zipfile.ZipFile(path) as archive:
            check(f"archive integrity: {name}", archive.testzip() is None and len(archive.namelist()) > 0)

docx_paths = sorted(FILES.rglob("*.docx"))
check("four support DOCX files plus the manuscript", len(docx_paths) == 5)
for path in docx_paths:
    doc = Document(path)
    text = "\n".join(p.text for p in doc.paragraphs)
    check(f"DOCX readable and titled: {path.name}", len(text.strip()) > 40)
    check(f"approved title present: {path.name}", TITLE in text)
    check(f"superseded title absent: {path.name}", "Economic change and the growth and organisation" not in text)

for relative in [
    "01_Manuscript/Paper1_Main_Manuscript.docx",
    "02_Title_Page/Paper1_Title_Page.docx",
    "05_Supplementary_Material/Paper1_Supplementary_Appendix.docx",
]:
    doc = Document(FILES / relative)
    text = "\n".join(p.text for p in doc.paragraphs)
    check(f"numbered author block: {Path(relative).name}", all(marker in text for marker in ("1. Qazi Raza, MPH", "2. Amna Khan, MPH", "3. Maham Qazi, BS Public Health")))
    check(f"corresponding-author label: {Path(relative).name}", "[Corresponding Author]" in text)

main_doc = Document(FILES / "01_Manuscript" / "Paper1_Main_Manuscript.docx")
main_text = "\n".join(p.text for p in main_doc.paragraphs)
for phrase in (
    "should be obtained",
    "should be inserted",
    "should review and confirm",
    "mandatory author-supplied",
):
    check(f"draft placeholder absent: {phrase}", phrase not in main_text.lower())
check(
    "final ethics declaration in package",
    "Ethical approval and individual informed consent were not required" in main_text,
)
check(
    "final author approval in package",
    "All authors read and approved the final manuscript" in main_text,
)

package_readme = (FILES / "README.md").read_text(encoding="utf-8")
check("submission package marked complete", "The manuscript and declarations are complete" in package_readme)

csv_tables = sorted((FILES / "03_Tables" / "CSV").glob("*.csv"))
check("five manuscript table CSV files", len(csv_tables) == 5)
check("three original PNG figures", len(list((FILES / "04_Figures").glob("Figure_[123]_*.png"))) == 3)
check("three 600 dpi TIFF figures", len(list((FILES / "04_Figures").glob("Figure_[123]_*_600dpi.tiff"))) == 3)
for path in (FILES / "04_Figures").glob("*.tiff"):
    with Image.open(path) as image:
        dpi = image.info.get("dpi", (0, 0))
        check(f"TIFF resolution metadata: {path.name}", dpi[0] >= 599 and dpi[1] >= 599)

bib = (FILES / "07_References" / "Paper1_References_45.bib").read_text(encoding="utf-8")
keys = re.findall(r"@[A-Za-z]+\{([^,]+),", bib)
check("45 unique BibTeX records", len(keys) == 45 and len(set(keys)) == 45)

manifest = FILES / "MANIFEST_SHA256.csv"
with manifest.open(encoding="utf-8") as handle:
    rows = list(csv.DictReader(handle))
check("complete package manifest populated", len(rows) >= 100)

failures = [label for label, passed in checks if not passed]
for label, passed in checks:
    print(f"{'PASS' if passed else 'FAIL'} | {label}")
print(f"SUMMARY | {len(checks) - len(failures)}/{len(checks)} checks passed")
sys.exit(1 if failures else 0)
