"""Validate the generated Paper 1 manuscript and its Word deliverable."""

from pathlib import Path
import re
import sys

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "manuscript" / "paper1_full_manuscript_vancouver.md"
DOCX = ROOT / "manuscript" / "paper1_full_manuscript_vancouver.docx"
TITLE = "Economic conditions and the changing scale and organisation of female sex work in Pakistan: a multi-period ecological study"

checks: list[tuple[str, bool]] = []


def check(name: str, condition: bool) -> None:
    checks.append((name, bool(condition)))


text = MD.read_text(encoding="utf-8")
doc = Document(DOCX)
doc_text = "\n".join(p.text for p in doc.paragraphs)

check("generated Markdown exists", MD.is_file())
check("generated Word manuscript exists", DOCX.is_file())
check("no unresolved citation keys", "[@" not in text)
check("no unresolved build markers", "[[" not in text and "]]" not in text)
check("approved title in Markdown", text.startswith(f"# {TITLE}\n"))
check("approved title in Word", TITLE in doc_text)
check("superseded title absent", "Economic change and the growth and organisation" not in text and "Economic change and the growth and organisation" not in doc_text)

required_sections = [
    "## Abstract",
    "## Introduction",
    "## Methods",
    "## Results",
    "## Discussion",
    "## Conclusion",
    "## Declarations",
    "## References",
]
for heading in required_sections:
    check(f"section present: {heading[3:]}", heading in text)

reference_block = text.split("## References", 1)[1]
reference_numbers = [int(n) for n in re.findall(r"(?m)^(\d+)\. ", reference_block)]
check("45 Vancouver references", reference_numbers == list(range(1, 46)))

body = text.split("## References", 1)[0]
cited_numbers: set[int] = set()
for content in re.findall(r"\[([0-9,\-]+)\]", body):
    for token in content.split(","):
        if "-" in token:
            start, end = (int(value) for value in token.split("-", 1))
            cited_numbers.update(range(start, end + 1))
        else:
            cited_numbers.add(int(token))
check("citation numbers are in range", bool(cited_numbers) and min(cited_numbers) >= 1 and max(cited_numbers) <= 45)
check("every bibliography entry is cited", cited_numbers == set(reference_numbers))

for exact in [
    "https://orcid.org/0000-0003-4303-3390",
    "https://orcid.org/0009-0000-9152-3912",
    "https://orcid.org/0009-0002-1199-7240",
    "qaziraza.rq@gmail.com",
    "amnaathar045@gmail.com",
    "qazimaham.mq@gmail.com",
]:
    check(f"author detail present: {exact}", exact in doc_text)

for exact in [
    f"1. Qazi Raza, MPH, Health Services Academy, NIH Complex, Park Road, Chak Shahzad, Islamabad, 44000, Pakistan [Corresponding Author]",
    f"2. Amna Khan, MPH, Health Services Academy, NIH Complex, Park Road, Chak Shahzad, Islamabad, 44000, Pakistan",
    f"3. Maham Qazi, BS Public Health, Health Services Academy, NIH Complex, Park Road, Chak Shahzad, Islamabad, 44000, Pakistan",
    "Email ID: qaziraza.rq@gmail.com",
    "Email ID: amnaathar045@gmail.com",
    "Email ID: qazimaham.mq@gmail.com",
]:
    check(f"approved author layout present: {exact}", exact in doc_text)

check("five manuscript tables", len(doc.tables) == 5)
check("three manuscript figures", len(doc.inline_shapes) == 3)
check("final ethics declaration present", "Ethical approval and individual informed consent were not required" in text)
check("final data-availability declaration present", "provided with this article in the accompanying supplementary and reproducibility files" in text)
check("final author approval present", "All authors read and approved the final manuscript" in text)
check("draft ethics placeholder absent", "should be obtained" not in text.lower())
check("repository placeholder absent", "should be inserted" not in text.lower())
check("author-confirmation placeholder absent", "should review and confirm" not in text.lower())
check("funding declaration present", "received no specific grant" in text)
check("competing interests declaration present", "declare no competing interests" in text)
check("AI-use disclosure present", "Generative artificial intelligence disclosure" in text)

for phrase in ["what was initially proposed", "working process", "our initial premise"]:
    check(f"process framing absent: {phrase}", phrase.lower() not in text.lower())

words = re.findall(r"\b[\w'-]+\b", body)
failures = [name for name, passed in checks if not passed]
for name, passed in checks:
    print(f"{'PASS' if passed else 'FAIL'} | {name}")
print(f"SUMMARY | {len(checks) - len(failures)}/{len(checks)} checks passed")
print(f"BODY WORDS | {len(words)}")
print(f"TABLES | {len(doc.tables)}")
print(f"FIGURES | {len(doc.inline_shapes)}")
print(f"REFERENCES | {len(reference_numbers)}")

sys.exit(1 if failures else 0)
